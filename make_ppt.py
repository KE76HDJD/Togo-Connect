"""Genere le rapport PowerPoint (10 slides) dans rapport/.

Design institutionnel : fond bleu nuit + accent vert #00693E,
cartes KPI encadrees, visuels avec legendes, tableau P1.
Lancer depuis la racine du projet : python3 make_ppt.py
"""
from pathlib import Path
import os
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

BASE = Path(__file__).resolve().parent
FIG = os.path.join(BASE, "outputs", "figures")
TAB = os.path.join(BASE, "data", "processed", "indicateurs_prefecture.csv")
OUT = os.path.join(BASE, "rapport", "Togo_Diagnostic_Connectivite_10slides.pptx")

NAVY = RGBColor(0x0B, 0x2C, 0x4D)
GREEN = RGBColor(0x00, 0x69, 0x3E)
BLUE = RGBColor(0x00, 0x66, 0xB3)
ORANGE = RGBColor(0xFF, 0x79, 0x00)
RED = RGBColor(0xD7, 0x26, 0x3D)
GREY = RGBColor(0x5B, 0x6B, 0x7B)
DARK = RGBColor(0x33, 0x47, 0x5B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF4, 0xF7, 0xFB)
FONT = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def solid_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, left, top, width, height, fill, line=None, radius=None):
    shape = (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    sp = slide.shapes.add_shape(shape, Inches(left), Inches(top), Inches(width), Inches(height))
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    if radius:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(1)
    return sp


def textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    box.text_frame.word_wrap = True
    return box


def para(tf, text, size, bold=False, color=DARK, align=PP_ALIGN.LEFT, first=False, space_after=4):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].text else tf.add_paragraph()
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = FONT
    p.alignment = align
    p.space_after = Pt(space_after)
    return p


def footer(slide, num):
    rect(slide, 0.6, 7.02, 12.13, 0.02, GREY)
    box = textbox(slide, 0.6, 7.08, 9.5, 0.35)
    para(box.text_frame, "ANKOU Yaokouma Kevin — Défi 1 Économie Numérique Togo", 10, False, GREY, first=True)
    box2 = textbox(slide, 11.4, 7.08, 1.33, 0.35)
    para(box2.text_frame, f"{num} / 10", 10, True, GREY, PP_ALIGN.RIGHT, first=True)


def content_head(slide, eyebrow, title, num):
    solid_bg(slide, WHITE)
    rect(slide, 0, 0, 13.333, 0.10, GREEN)
    box = textbox(slide, 0.7, 0.35, 11.9, 0.35)
    para(box.text_frame, eyebrow, 11, True, GREEN, first=True)
    tbox = textbox(slide, 0.7, 0.68, 11.9, 0.7)
    para(tbox.text_frame, title, 27, True, NAVY, first=True)
    rect(slide, 0.7, 1.42, 2.4, 0.055, GREEN)
    footer(slide, num)


def bullets(slide, left, top, width, items, size=15):
    box = textbox(slide, left, top, width, 5.2)
    tf = box.text_frame
    for i, (head, body) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        r1 = p.add_run()
        r1.text = "▪ " + head + "  "
        r1.font.size = Pt(size)
        r1.font.bold = True
        r1.font.color.rgb = NAVY
        r1.font.name = FONT
        r2 = p.add_run()
        r2.text = body
        r2.font.size = Pt(size)
        r2.font.color.rgb = DARK
        r2.font.name = FONT
    return box


def framed_image(slide, name, left, top, width, caption):
    """Visuel encadré sur fond clair + légende (jamais de chevauchement)."""
    img_h = 3.9
    rect(slide, left - 0.12, top - 0.12, width + 0.24, img_h + 0.75, LIGHT, line=RGBColor(0xE3, 0xE9, 0xEF), radius=0.08)
    p = os.path.join(FIG, name)
    if os.path.exists(p):
        slide.shapes.add_picture(p, Inches(left), Inches(top), width=Inches(width), height=Inches(img_h))
    cap = textbox(slide, left, top + img_h + 0.05, width, 0.5)
    para(cap.text_frame, caption, 11, False, GREY, PP_ALIGN.CENTER, first=True)


def kpi_card(slide, left, top, width, value, label, color):
    card = rect(slide, left, top, width, 1.55, WHITE, radius=0.12)
    card.shadow.inherit_shadow = False
    tf = textbox(slide, left + 0.15, top + 0.12, width - 0.3, 1.35).text_frame
    para(tf, value, 27, True, color, PP_ALIGN.CENTER, first=True, space_after=1)
    para(tf, label, 11.5, False, DARK, PP_ALIGN.CENTER)


def info_card(slide, left, top, width, height, title, lines, accent):
    rect(slide, left, top, width, height, WHITE, line=RGBColor(0xE3, 0xE9, 0xEF), radius=0.1)
    rect(slide, left, top, 0.09, height, accent)
    box = textbox(slide, left + 0.3, top + 0.15, width - 0.5, height - 0.3)
    tf = box.text_frame
    para(tf, title, 14, True, NAVY, first=True, space_after=4)
    for ln in lines:
        para(tf, "•  " + ln, 12.5, False, DARK, space_after=3)


# ============ 1 - COUVERTURE ============
s = prs.slides.add_slide(BLANK)
solid_bg(s, NAVY)
rect(s, 0, 0, 0.32, 7.5, GREEN)
rect(s, 0.32, 0, 13.013, 0.07, GREEN)
box = textbox(s, 0.9, 0.55, 11.5, 0.4)
para(box.text_frame, "TOGO — ÉCONOMIE NUMÉRIQUE  •  DÉFI 1", 15, True, RGBColor(0x7F, 0xDB, 0xA6), first=True)
tbox = textbox(s, 0.9, 1.0, 11.5, 2.2)
tf = tbox.text_frame
para(tf, "Diagnostic de l'accès aux", 38, True, WHITE, first=True, space_after=0)
para(tf, "télécommunications et services numériques", 38, True, WHITE, space_after=6)
para(tf, "Cartographie des infrastructures, analyse du Mobile Money et priorités d'extension.", 16, False, RGBColor(0xC9, 0xD6, 0xE3))
for i, (val, lab, col) in enumerate([
        ("90", "Agences\nMoov 28 • Togocom 62", BLUE),
        ("19 788", "Points Mobile Money\n39 préfectures", ORANGE),
        ("3", "Datacenters\ntous à Lomé", GREEN),
        ("12 / 39", "Préfectures sans agence\nZones P1", RED)]):
    left = 0.9 + i * 3.0
    card = rect(s, left, 4.15, 2.75, 1.7, WHITE, radius=0.12)
    tfc = textbox(s, left + 0.15, 4.25, 2.45, 1.5).text_frame
    para(tfc, val, 30, True, col, PP_ALIGN.CENTER, first=True, space_after=1)
    for j, ln in enumerate(lab.split("\n")):
        para(tfc, ln, 11.5, j == 0, NAVY if j == 0 else GREY, PP_ALIGN.CENTER, space_after=0)
abox = textbox(s, 0.9, 6.25, 11.5, 0.9)
para(abox.text_frame, "Préparé par ANKOU Yaokouma Kevin  •  Dashboard Python + Rapport  •  Données ouvertes geodata.gouv.tg", 13, False, RGBColor(0xC9, 0xD6, 0xE3), first=True)
para(abox.text_frame, "Septembre 2026 — Échelle Région / Préfecture", 12, False, RGBColor(0x8F, 0xA3, 0xBF))

# ============ 2 - PROBLÉMATIQUE ============
s = prs.slides.add_slide(BLANK)
content_head(s, "DIAGNOSTIC CONNECTIVITÉ TOGO", "Problématique et objectifs", 2)
bullets(s, 0.7, 1.8, 11.9, [
    ("Contexte.", "L'infrastructure visible sur une carte ne dit pas qui accède réellement aux services numériques."),
    ("Question centrale.", "Comment les agences, les datacenters et les points Mobile Money se répartissent-ils face aux populations ?"),
    ("Objectifs.", "Cartographier l'offre, mesurer l'adéquation Mobile Money, croiser offre et territoires, détecter les zones blanches, recommander."),
    ("Périmètre réel.", "90 agences, 19 788 points Mobile Money, 3 datacenters, fichier CANAL+ vide — Région / Préfecture."),
])

# ============ 3 - MÉTHODOLOGIE ============
s = prs.slides.add_slide(BLANK)
content_head(s, "DIAGNOSTIC CONNECTIVITÉ TOGO", "Méthodologie et données", 3)
info_card(s, 0.7, 1.8, 3.75, 3.6, "1. Sources",
          ["6 fichiers geodata.gouv.tg", "Agences, Mobile Money,", "datacenters + géométries"], GREEN)
info_card(s, 4.75, 1.8, 3.75, 3.6, "2. Traitement",
          ["Télécom = union dédupliquée", "Nsp vers Inconnu", "Grand Lomé = Golfe + Agoè-Nyivé"], BLUE)
info_card(s, 8.8, 1.8, 3.75, 3.6, "3. Indicateurs",
          ["Points MoMo par agence", "Priorités P1 / P2 / P3", "Ratio national 220 / 1"], ORANGE)
lim = textbox(s, 0.7, 5.65, 11.9, 1.1)
para(lim.text_frame, "Limites assumées : pas de population RGPH ni d'antennes BTS fournies — zones blanches relatives, version 2 quantifiée en points pour 10 000 habitants.",
     13, False, GREY, first=True)

# ============ 4 - OFFRE AGENCES ============
s = prs.slides.add_slide(BLANK)
content_head(s, "DIAGNOSTIC CONNECTIVITÉ TOGO", "Offre : des agences concentrées à Lomé", 4)
bullets(s, 0.7, 1.8, 5.6, [
    ("Fracture territoriale.", "Maritime 51/90 (57 %) contre Savanes 6, Centrale 11, Kara 10, Plateaux 12."),
    ("Deux stratégies.", "Togocom (62) maille 27 préfectures ; Moov (28) couvre 15 préfectures."),
    ("Point aveugle.", "CANAL+ : fichier vide — segment à publier en open data."),
])
framed_image(s, "01_agences_region.png", 6.9, 1.9, 5.5, "Figure 1 — Agences par région (n = 90)")

# ============ 5 - DATACENTERS ============
s = prs.slides.add_slide(BLANK)
content_head(s, "DIAGNOSTIC CONNECTIVITÉ TOGO", "Datacenters : un risque de concentration", 5)
bullets(s, 0.7, 1.8, 5.6, [
    ("3 sites, 1 commune.", "Lomé Data Center (2022), E-Gouv NOC (2016), Café Informatique Cloud (2021) — tous à Golfe."),
    ("Vulnérabilité.", "Aucune redondance intérieure : panne unique et latence vers le nord."),
    ("Recommandation.", "Nœud secondaire à Kara ou Sokodé adossé à la fibre."),
])
framed_image(s, "06_datacenters.png", 6.9, 1.9, 5.5, "Figure 2 — Datacenters par préfecture (n = 3)")

# ============ 6 - MOBILE MONEY ============
s = prs.slides.add_slide(BLANK)
content_head(s, "DIAGNOSTIC CONNECTIVITÉ TOGO", "Mobile Money : le vrai maillage du pays", 6)
bullets(s, 0.7, 1.8, 5.6, [
    ("Couverture totale.", "19 788 points sur 39 préfectures : Golfe 5 121 contre Mô 32."),
    ("Interopérabilité de fait.", "64 % de points mixtes Moov et Togocom ; 24 % Togocom seul."),
    ("Levier n° 1.", "Les 12 zones sans agence sont déjà desservies en Mobile Money."),
])
framed_image(s, "02_momo_region.png", 6.9, 1.9, 5.5, "Figure 3 — Points Mobile Money par région (n = 19 788)")

# ============ 7 - ADÉQUATION ============
s = prs.slides.add_slide(BLANK)
content_head(s, "DIAGNOSTIC CONNECTIVITÉ TOGO", "Adéquation : la tension est hors Lomé", 7)
bullets(s, 0.7, 1.8, 5.6, [
    ("Ratio national.", "220 points Mobile Money pour 1 agence."),
    ("Tension maximale.", "Lacs 741/0, Vo 476/0, Wawa 244/0 : forte demande sans relais."),
    ("Lecture du graphique.", "En haut à gauche : beaucoup de Mobile Money, aucune agence."),
])
framed_image(s, "04_adequation.png", 6.9, 1.9, 5.5, "Figure 4 — Chaque point = une préfecture")

# ============ 8 - ZONES BLANCHES (TABLEAU) ============
s = prs.slides.add_slide(BLANK)
content_head(s, "DIAGNOSTIC CONNECTIVITÉ TOGO", "Zones blanches relatives : 12 préfectures P1", 8)
ind = pd.read_csv(TAB).query("priorite == 1").sort_values("nb_momo", ascending=False)
rows, cols = len(ind) + 1, 4
tbl_shape = s.shapes.add_table(rows, cols, Inches(0.7), Inches(1.75), Inches(7.6), Inches(4.9))
tbl = tbl_shape.table
widths = [3.0, 1.9, 1.5, 1.2]
for j, w in enumerate(widths):
    tbl.columns[j].width = Inches(w)
headers = ["Préfecture", "Région", "Points MoMo", "Agences"]
for j, h in enumerate(headers):
    c = tbl.cell(0, j)
    c.text_frame.paragraphs[0].text = h
    for p in c.text_frame.paragraphs:
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.font.name = FONT
    c.fill.solid()
    c.fill.fore_color.rgb = NAVY
for i, r in enumerate(ind.itertuples(), start=1):
    for j, val in enumerate([r.prefecture, r.region, str(r.nb_momo), str(r.nb_agences)]):
        c = tbl.cell(i, j)
        c.text_frame.paragraphs[0].text = val
        for p in c.text_frame.paragraphs:
            p.font.size = Pt(11)
            p.font.color.rgb = DARK
            p.font.name = FONT
        if i % 2 == 0:
            c.fill.solid()
            c.fill.fore_color.rgb = LIGHT
info_card(s, 8.7, 1.75, 3.9, 4.9, "Clé de lecture",
          ["P1 : 0 agence + Mobile Money", "isolé — à équiper en premier", "P2 : agence unique fragile", "P3 : territoires desservis", "Couverture radio à valider", "par mesures ARCEP 2G/3G/4G"], RED)

# ============ 9 - RECOS R1-R2 ============
s = prs.slides.add_slide(BLANK)
content_head(s, "DIAGNOSTIC CONNECTIVITÉ TOGO", "Recommandations R1-R2 : gains rapides", 9)
info_card(s, 0.7, 1.8, 5.6, 2.3, "R1 — Agences mobiles mutualisées (0-12 mois, coût faible)",
          ["Bus Moov + Togocom sur les 12 P1", "Indicateur : 12 tournées mensuelles"], GREEN)
info_card(s, 0.7, 4.25, 5.6, 2.3, "R2 — 200 relais Mobile Money labellisés (6-18 mois)",
          ["Formation + kit de services", "Indicateur : +40 % d'actes hors Lomé"], BLUE)
framed_image(s, "03_top_momo.png", 6.9, 1.9, 5.5, "Figure 5 — Cibles R2 : top Mobile Money sans agence")

# ============ 10 - RECOS R3-R4 + FEUILLE DE ROUTE ============
s = prs.slides.add_slide(BLANK)
content_head(s, "DIAGNOSTIC CONNECTIVITÉ TOGO", "Recommandations R3-R4 et feuille de route", 10)
info_card(s, 0.7, 1.8, 5.9, 2.2, "R3 — Datacenter Kara / Sokodé (12-36 mois)",
          ["Redondance + fibre, reprise < 4 h"], GREEN)
info_card(s, 6.75, 1.8, 5.9, 2.2, "R4 — Transparence open data (3 mois, coût nul)",
          ["CANAL+, population RGPH, antennes BTS"], BLUE)
steps = [("T0", "R4\nPublier"), ("T+6", "R1\nBus mobiles"), ("T+12", "R2\n200 relais"), ("T+24", "R3\nDatacenter")]
for i, (t, lab) in enumerate(steps):
    left = 0.7 + i * 3.1
    rect(s, left, 4.25, 2.85, 1.5, LIGHT, line=GREEN, radius=0.1)
    tf = textbox(s, left + 0.15, 4.3, 2.55, 1.4).text_frame
    para(tf, t, 20, True, GREEN, PP_ALIGN.CENTER, first=True, space_after=1)
    for ln in lab.split("\n"):
        para(tf, ln, 12.5, True, NAVY, PP_ALIGN.CENTER, space_after=0)
box = textbox(s, 0.7, 5.95, 11.9, 0.9)
para(box.text_frame, "Pilotes : ARCEP, Togo Digital, opérateurs — Suivi : file d'attente, actes hors Lomé, latence nord, jeux publiés.",
     12.5, False, GREY, first=True)

prs.save(OUT)
print(f"PPT OK: {OUT} ({len(prs.slides)} slides)")
